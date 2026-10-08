"""Koncept: automatické třídění podkladů k pojistné události (havárie auta).

Aplikace hlídá složku klient_upload/. Jakmile tam klient vloží soubor:
  1. přečte jeho text (PDF, Word, Excel, e-mail, HTML, TXT; obrázky a skeny přes OCR),
  2. pozná typ dokumentu (faktura, policejní protokol, fotka, ...),
  3. vytáhne z dokumentu důležité údaje (SPZ, VIN, částky, data, ...),
  4. přesune soubor do roztridene/<typ>/,
  5. přepíše jeden souhrnný soubor roztridene/SOUHRN_pojistne_udalosti.html.

Spuštění:  .venv/bin/python zpracuj_dokumenty.py           (hlídá složku, ukončení Ctrl+C)
           .venv/bin/python zpracuj_dokumenty.py --jednou  (zpracuje, co tam je, a skončí)
"""
import email
import hashlib
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from collections import Counter
from datetime import datetime
from email import policy
from html.parser import HTMLParser
from pathlib import Path

BASE = Path(__file__).parent
INBOX = BASE / "klient_upload"
OUT = BASE / "roztridene"
DB = OUT / ".databaze.json"
SUMMARY = OUT / "SOUHRN_pojistne_udalosti.html"

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
IGNORE_SUFFIX = {".part", ".crdownload", ".tmp", ".swp"}

# Typ dokumentu -> (název složky, klíčová slova s váhou) – používá se, jen když není Claude.. Slova jsou bez diakritiky,
# porovnávají se s textem bez diakritiky a bez mezer (OCR mezery často slije nebo vynechá).
TYPES = {
    "hlaseni": ("01_Hlaseni_skody", {"hlasim": 4, "hlaseniskody": 3, "hlasenipojistneudalosti": 4,
                                     "skodnouudalost": 2, "spozdravem": 1, "prosimo": 1}),
    "zaznam": ("02_Zaznam_o_nehode", {"zaznamodopravninehode": 5, "spolecnyzaznam": 2, "vozidloa": 2,
                                      "vozidlob": 2, "zelenekarty": 2, "podpisyridicu": 2}),
    "protokol": ("03_Policejni_protokol", {"policie": 2, "protokolonehode": 5, "c.j.": 3,
                                           "dopravniinspektorat": 3, "zavineni": 3, "sepsal": 1}),
    "doklad_vozidla": ("04_Doklady_vozidla", {"osvedcenioregistraci": 5, "technickyprukaz": 5,
                                              "datumprvniregistrace": 3, "provozovatel": 2,
                                              "tovarniznacka": 2, "zdvihovyobjem": 2}),
    "kalkulace": ("05_Kalkulace_opravy", {"kalkulace": 4, "odhadskody": 2, "rozpocetopravy": 4,
                                          "zakazka": 1}),
    "faktura": ("06_Faktury", {"faktura": 3, "danovydoklad": 3, "splatnost": 2, "celkemkuhrade": 2,
                                "variabilnisymbol": 2, "prijmovydoklad": 3, "prijmovypokladnidoklad": 4,
                                "rechnungsnummer": 5, "rechnung": 4, "betrag": 2, "castka": 2,
                                "kuhrade": 2, "doklad": 1, "parkoviste": 1,
                                "dodavatel": 1, "odberatel": 1}),
    "lekarska_zprava": ("10_Lekarske_zpravy", {"lekarskazprava": 5, "ambulantnizprava": 5,
                                                   "pracovnineschopnost": 3, "pacientudava": 2}),
    "korespondence": ("11_Korespondence", {"skodniudalosti": 3, "omlouvam": 2,
                                            "volejte": 2, "sms": 3, "dobryden": 1,
                                            "poznamka": 2, "telefonicky": 2}),
}
# Složky a názvy pro všechny typy. Pravidla umí jen typy z TYPES, Claude rozezná všechny.
FOLDER = {k: v[0] for k, v in TYPES.items()} | {
    "foto": "07_Fotodokumentace", "ridicsky_prukaz": "08_Ridicsky_prukaz",
    "pojistna_smlouva": "09_Pojistna_smlouva", "lekarska_zprava": "10_Lekarske_zpravy",
    "korespondence": "11_Korespondence", "jiny_k_udalosti": "12_Ostatni_k_udalosti",
    "neznamy": "99_Nerozpoznano"}
TYPE_LABEL = {"hlaseni": "Hlášení škody", "zaznam": "Záznam o dopravní nehodě",
              "protokol": "Policejní protokol", "doklad_vozidla": "Doklad vozidla (OR/TP)",
              "ridicsky_prukaz": "Řidičský průkaz", "pojistna_smlouva": "Pojistná smlouva / zelená karta",
              "kalkulace": "Kalkulace opravy", "faktura": "Faktura / účtenka",
              "lekarska_zprava": "Lékařská zpráva", "korespondence": "Korespondence (SMS, chat, dopis)",
              "jiny_k_udalosti": "Jiný dokument k události", "foto": "Fotodokumentace",
              "neznamy": "Nerozpoznaný / nesouvisející dokument"}


# ----------------------------------------------------------------- text bez diakritiky
def ascii_same_length(text: str) -> str:
    """'Škoda Ř' -> 'Skoda R'. Každý znak zůstane na stejné pozici, takže co najdeme
    v textu bez diakritiky, můžeme vyříznout z originálu i s háčky."""
    return "".join(unicodedata.normalize("NFD", c)[0] for c in text)


def compact(text: str) -> str:
    return re.sub(r"\s+", "", ascii_same_length(text).lower())


# ----------------------------------------------------------------- čtení textu
_ocr = None


def ocr_image(path: Path) -> str:
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR()
    result, _ = _ocr(str(path))
    if not result:
        return ""
    # OCR vrací jednotlivé kousky textu; poskládáme je zpět do řádků podle výšky na stránce.
    boxes = sorted(((sum(p[1] for p in b) / 4, b[0][0], abs(b[3][1] - b[0][1]), t) for b, t, _ in result))
    lines, cur, cur_y = [], [], None
    for y, x, h, t in boxes:
        if cur_y is not None and abs(y - cur_y) > max(h, 10) * 0.6:
            lines.append(" ".join(t for _, t in sorted(cur)))
            cur = []
        cur.append((x, t))
    if cur:
        lines.append(" ".join(t for _, t in sorted(cur)))
    return "\n".join(lines)


def read_pdf(path: Path) -> tuple[str, bool]:
    from pypdf import PdfReader
    text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if len(text.strip()) > 30:
        return text, False
    # Naskenované PDF: uvnitř je jen obrázek -> převést stránky na obrázky a přečíst OCR.
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["pdftoppm", "-r", "200", "-png", str(path), f"{tmp}/str"], check=True)
        return "\n".join(ocr_image(p) for p in sorted(Path(tmp).glob("*.png"))), True


def read_docx(path: Path) -> str:
    from docx import Document
    d = Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:  # sloučené buňky se v python-docx opakují
                if not cells or cells[-1] != c.text.strip():
                    cells.append(c.text.strip())
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path)  # se vzorci, jednoduché SUM spočítáme sami
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            vals = []
            for c in row:
                v = c.value
                if isinstance(v, str) and (m := re.fullmatch(r"=SUM\((\w+\d+):(\w+\d+)\)", v.upper())):
                    v = sum(x.value for r in ws[m[1]:m[2]] for x in r if isinstance(x.value, (int, float)))
                if v is not None:
                    vals.append(f"{v:.2f}".replace(".", ",") if isinstance(v, float) else str(v))
            if vals:
                lines.append(" | ".join(vals))
    return "\n".join(lines)


def read_eml(path: Path) -> str:
    msg = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    body = msg.get_body(preferencelist=("plain", "html"))
    content = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        content = strip_html(content)
    return f"Od: {msg['From']}\nKomu: {msg['To']}\nPředmět: {msg['Subject']}\nDatum: {msg['Date']}\n\n{content}"


class _TextOnly(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self.skip = [], False

    def handle_starttag(self, tag, attrs):
        self.skip = tag in ("style", "script", "title")
        if tag in ("p", "br", "tr", "h1", "h2", "h3", "div", "li"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        self.skip = False

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def strip_html(source: str) -> str:
    p = _TextOnly()
    p.feed(source)
    return re.sub(r"[ \t]+", " ", "".join(p.parts))


def read_text(path: Path) -> tuple[str, str]:
    """Vrátí (text, způsob čtení)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        text, scanned = read_pdf(path)
        return text, "OCR naskenovaného PDF" if scanned else "text z PDF"
    if ext == ".docx":
        return read_docx(path), "Word"
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path), "Excel"
    if ext == ".eml":
        return read_eml(path), "e-mail"
    if ext in (".html", ".htm"):
        return strip_html(path.read_text(encoding="utf-8", errors="replace")), "HTML"
    if ext in (".txt", ".csv", ".md"):
        return path.read_text(encoding="utf-8", errors="replace"), "text"
    if ext in IMAGE_EXT:
        return ocr_image(path), "OCR obrázku"
    return "", "nepodporovaný formát"


# ----------------------------------------------------------------- rozpoznání typu
def classify(text: str, is_image: bool) -> tuple[str, dict]:
    flat = compact(text)
    scores = {t: sum(w for kw, w in kws.items() if kw in flat) for t, (_, kws) in TYPES.items()}
    best = max(scores, key=scores.get)
    letters = len(re.findall(r"[A-Za-z0-9]", ascii_same_length(text)))

    # Pokud je obrázek typu „vyfocený doklad“, neměli bychom ho hned hodit do fotodokumentace,
    # jen protože OCR zachytil málo textu. Takový obrázek má často jen pár slov, ale zřetelné
    # klíčové slovo typu "faktura", "kalkulace", "protokol" apod. -> preferujeme dokument.
    if is_image and letters < 60:
        if scores[best] >= 3:
            return best, scores
        return "foto", scores
    if scores[best] < 3:
        return "neznamy", scores
    return best, scores


# ----------------------------------------------------------------- vytažení údajů
AMOUNT = r"(\d{1,3}(?:[  ]\d{3})+(?:[,.]\d{2})?|\d+(?:[,.]\d{1,2})?)"
DATE = r"(\d{1,2}\.\s?\d{1,2}\.\s?\d{4})"


def find(text: str, pattern: str, flags=re.I) -> str | None:
    """Hledá v textu bez diakritiky, ale vrací kus originálu (i s háčky).
    Mezera ve vzoru znamená „libovolně mezer, i žádná“ – OCR slova často slije."""
    pattern = pattern.replace(" ", r"\s*")
    m = re.search(pattern, ascii_same_length(text), flags | re.M)
    if not m:
        return None
    return text[m.start(1):m.end(1)].strip(" :|,")


def norm_date(s: str | None) -> str | None:
    if not s:
        return None
    d, m, y = re.findall(r"\d+", s)[:3]
    return f"{int(d):02d}.{int(m):02d}.{y}"


def to_number(s: str | None) -> float | None:
    if not s:
        return None
    s = s.replace(" ", "").replace(" ", "")
    s = s.replace(",", ".") if "," in s else s
    try:
        return float(s)
    except ValueError:
        return None


def czk(v: float | None) -> str:
    return "" if v is None else f"{v:,.2f} Kč".replace(",", " ").replace(".", ",")


def plates(text: str) -> list[str]:
    found = re.findall(r"(?:(?<=SPZ)|(?<=RZ)|(?<![A-Z0-9]))(\d[A-Z][A-Z0-9])\s?(\d{4})\b", ascii_same_length(text).upper())
    return list(dict.fromkeys(f"{a} {b}" for a, b in found))


def vins(text: str) -> list[str]:
    found = re.findall(r"\b[A-HJ-NPR-Z0-9]{17}\b", ascii_same_length(text).upper())
    return list(dict.fromkeys(v for v in found if re.search(r"\d", v) and re.search(r"[A-Z]", v)))


def extract(doc_type: str, text: str, filename: str) -> dict:
    f: dict = {}
    if doc_type == "faktura":
        document_number = (find(text, r"(?:cislo faktury|rechnungsnummer|doklad\.?\s*(?:c\.?)?)\s*:?\s*([A-Z0-9/\-]+)")
                           or find(text, r"faktura\s*c?\.?\s*:?\s*([A-Z0-9/\-]*\d[A-Z0-9/\-]*)"))
        f["Číslo faktury"] = document_number if document_number and re.search(r"\d", document_number) else None
        supplier = (find(text, r"dodavatel\s*:?\s*([^\n,|]+)")
                    or find(text, r"^\s*(?:(?:rechnung|faktura)\s*)?([A-Z][^\n,|]*?(?:gmbh|s\.r\.o\.|a\.s\.))")
                    or find(text, r"\b(AUTOSKLO[A-Z0-9]+)\b")
                    or find(text, r"^(Parkoviste\s+[^,\n]+)"))
        if not supplier:
            first_line = text.strip().splitlines()[0].strip() if text.strip() else ""
            supplier = first_line if len(first_line) <= 100 else None
        f["Dodavatel"] = supplier
        f["IČO dodavatele"] = find(text, r"\b[I1]C[O0]?\s*:?\s*(\d{8})")
        f["Datum vystavení"] = norm_date(find(text, r"(?:datum vystaveni|rechnungsdatum|vystaveno|dne)\s*:?\s*" + DATE))
        f["Splatnost"] = norm_date(find(text, r"(?:datum splatnosti|splatnost)\s*:?\s*" + DATE))
        total = (find(text, r"(?:celkem\s*k\s*uhrade|kuhrade|gesamtbetrag|summe)\s*:?\s*" + AMOUNT)
             or find(text, r"(?:castka|betrag)\s*:?\s*" + AMOUNT))
        f["_castka"] = to_number(total)
        flat = compact(text)
        f["Měna"] = "EUR" if any(marker in flat for marker in ("eur", "osterreich", "oesterreich")) or "\x08" in text else "CZK"
        display_currency = "Kč" if f["Měna"] == "CZK" else f["Měna"]
        f["Celkem k úhradě"] = ("" if f["_castka"] is None else
                    f"{f['_castka']:,.2f} {display_currency}".replace(",", " ").replace(".", ","))
        f["Za co"] = ("výměna čelního skla" if "celni" in flat and "skl" in flat else
                      "oprava dveří" if "dver" in flat and any(k in flat for k in ("oprava", "oprav")) else
                      "výměna světlometu" if any(k in flat for k in ("scheinwerfer", "headlight", "svetlomet")) else
                  "parkování vozidla" if "parkovist" in flat else
                  "odtah vozidla" if "odtah" in flat else
                      "náhradní vozidlo" if "nahradnivozidlo" in flat or "pujcovna" in flat else
                      "oprava vozidla" if any(k in flat for k in ("prace", "lakovani", "naraznik")) else "jiné")
    elif doc_type == "protokol":
        f["Číslo jednací"] = find(text, r"c\.\s*j\.\s*:?\s*(\S+)")
        f["Datum nehody"] = norm_date(find(text, r"datum a cas nehody\s*:?\s*" + DATE))
        f["Čas nehody"] = find(text, r"datum a cas nehody.*?(\d{1,2}:\d{2})")
        f["Místo nehody"] = find(text, r"misto nehody\s*:?\s*(.+)")
        f["Účastník 1"] = find(text, r"ucastnik 1\s*:?\s*([^,\n]+)")
        f["Účastník 2"] = find(text, r"ucastnik 2\s*:?\s*([^,\n]+)")
        f["Zavinění"] = find(text, r"zavineni\s*:?\s*(.+)")
        f["Zranění"] = find(text, r"zraneni osob\s*:?\s*([^.\n]+)")
        f["Odhad škody (policie)"] = czk(to_number(find(text, r"odhad hmotne skody\s*:?\s*" + AMOUNT)))
    elif doc_type == "zaznam":
        f["Datum nehody"] = norm_date(find(text, r"datum nehody\s*:?\s*" + DATE))
        f["Čas nehody"] = find(text, r"cas\s*:?\s*(\d{1,2}:\d{2})")
        f["Místo nehody"] = find(text, r"misto\s*:?\s*(.+)")
        for label, key in (("Řidič", "ridic"), ("Pojistitel", "pojistitel"),
                           ("Pojistná smlouva", "cislo pojistne smlouvy"), ("Poškození", "viditelna poskozeni")):
            row = find(text, rf"^{key}\s*\|\s*(.+)$")
            if row:
                a, _, b = row.partition("|")
                f[f"{label} – vozidlo A"], f[f"{label} – vozidlo B"] = a.strip(), b.strip()
        f["Okolnosti"] = find(text, r"okolnosti\s*:?\s*(.+)")
    elif doc_type == "hlaseni":
        f["Odesílatel"] = find(text, r"^od\s*:\s*(.+)")
        f["Předmět"] = find(text, r"^predmet\s*:\s*(.+)")
        f["Přijato"] = find(text, r"^datum\s*:\s*(.+)")
        f["Datum nehody"] = norm_date(find(text, r"dne\s+" + DATE))
        f["Pojistná smlouva"] = find(text, r"smlouvy\s*:?\s*(\d{6,12})")
        f["Telefon"] = find(text, r"tel\.?\s*:?\s*(\+?[\d ]{9,16}\d)")
        f["Požadavky klienta"] = find(text, r"(prosim[^\n]+)")
    elif doc_type == "lekarska_zprava":
        f["Pacient"] = find(text, r"pacient\s*:?\s*([^,\n]+)")
        f["Datum vyšetření"] = norm_date(find(text, r"datum vysetreni\s*:?\s*" + DATE))
        f["Diagnóza"] = find(text, r"dg\.?\s*:?\s*([^\n]+)")
        f["Léčba"] = find(text, r"th\.?\s*:?\s*([^\n]+)")
    elif doc_type == "korespondence":
        f["Číslo škody u druhé pojišťovny"] = find(text, r"cislo skodni udalosti u nich je\s*([^\n.]+)")
        f["Telefon"] = find(text, r"(?:volejte|tel\.?)\s*:?[ ]*(\+?[\d ]{9,16}\d)")
    elif doc_type == "doklad_vozidla":
        f["Registrační značka"] = find(text, r"registracni znacka\s*:?\s*(\w{3}\s?\d{4})")
        f["VIN"] = find(text, r"\(VIN\)\s*:?\s*([A-Z0-9]{17})")
        f["Provozovatel"] = find(text, r"provozovatel\s*:?\s*(.+)")
        f["Značka a model"] = " ".join(filter(None, (find(text, r"tovarni znacka\s*:?\s*(.+)"),
                                                     find(text, r"obchodni oznaceni\s*:?\s*(.+)"))))
        f["První registrace"] = norm_date(find(text, r"prvni registrace\s*:?\s*" + DATE))
        f["Číslo osvědčení"] = find(text, r"cislo osvedceni\s*:?\s*(.+)")
    elif doc_type == "kalkulace":
        f["Zakázka"] = find(text, r"zakazka\s*:?\s*\|?\s*(\S+)")
        f["Datum kalkulace"] = norm_date(find(text, r"datum kalkulace\s*:?\s*\|?\s*" + DATE))
        f["Zpracoval"] = find(text, r"zpracoval\s*:?\s*\|?\s*(.+)")
        bez = to_number(find(text, r"celkem bez dph\s*[:\|]?\s*" + AMOUNT))
        f["Celkem bez DPH"] = czk(bez)
        f["_castka"] = to_number(find(text, r"celkem s dph\s*[:\|]?\s*" + AMOUNT))
        f["Celkem s DPH"] = czk(f["_castka"])
    elif doc_type == "foto":
        m = re.search(r"(20\d{2})(\d{2})(\d{2})_?(\d{2})?(\d{2})?", filename)
        if m:
            f["Pořízeno (dle názvu)"] = f"{m[3]}.{m[2]}.{m[1]}" + (f" {m[4]}:{m[5]}" if m[4] else "")
    # Společné údaje, které hledáme ve všech dokumentech
    f["_spz"] = plates(text)
    f["_vin"] = vins(text)
    f["_pu"] = list(dict.fromkeys(re.findall(r"PU-\d{4}-\d+", text)))
    return {k: v for k, v in f.items() if v not in (None, "", [])}


CSS = """
:root{--bg:#f6f7f9;--card:#fff;--text:#1d2330;--muted:#5b6475;--line:#e2e5eb;--accent:#1f5fa8;
--ok:#1e7b3c;--okbg:#e6f4ea;--warn:#9a4b00;--warnbg:#fff1e0}
@media (prefers-color-scheme:dark){:root{--bg:#14171c;--card:#1d2128;--text:#e6e8ec;--muted:#a3aab8;
--line:#2e333d;--accent:#7fb2ee;--ok:#7fd49a;--okbg:#17301f;--warn:#ffb866;--warnbg:#3a2a14}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);
font:15px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1000px;margin:0 auto;padding:24px 16px 60px}
h1{font-size:26px;margin:0 0 4px}h2{font-size:19px;margin:32px 0 10px}
.muted{color:var(--muted)}.card{background:var(--card);border:1px solid var(--line);border-radius:10px;
padding:14px 16px;margin:10px 0}
table{border-collapse:collapse;width:100%}td,th{text-align:left;padding:6px 8px;border-bottom:1px solid var(--line);
vertical-align:top}th{color:var(--muted);font-weight:600;width:34%}
.num{text-align:right;white-space:nowrap}.check{padding:6px 10px;border-radius:6px;margin:4px 0}
.ok{background:var(--okbg);color:var(--ok)}.warn{background:var(--warnbg);color:var(--warn)}
.tag{display:inline-block;font-size:12px;padding:1px 8px;border-radius:99px;background:var(--line);margin-left:6px}
a{color:var(--accent)}details summary{cursor:pointer;color:var(--muted)}
pre{white-space:pre-wrap;font-size:12.5px;background:var(--bg);padding:10px;border-radius:6px;overflow-x:auto}
.photos{display:flex;flex-wrap:wrap;gap:12px}figure{margin:0;width:220px;max-width:100%;font-size:12.5px;color:var(--muted)}.photos img{width:100%;border-radius:8px;
border:1px solid var(--line)}.src{font-size:12px;color:var(--muted)}
"""


# ----------------------------------------------------------------- údaje o události z pravidel
def rule_fakta(t: str, u: dict) -> dict:
    """Převede údaje vytažené pravidly na stejná „fakta“, jaká vrací Claude."""
    spz = u.get("_spz", [])
    f = {
        "cislo_pojistne_udalosti": (u.get("_pu") or [None])[0],
        "vin": (u.get("_vin") or [None])[0],
        "spz_klienta": spz[0] if spz else None,
        "spz_druheho_vozidla": spz[1] if t in ("zaznam", "protokol", "hlaseni") and len(spz) > 1 else None,
        "datum_nehody": u.get("Datum nehody"),
        "cas_nehody": u.get("Čas nehody"),
        "misto_nehody": u.get("Místo nehody"),
        "klient": u.get("Řidič – vozidlo A") or u.get("Účastník 1") or u.get("Provozovatel"),
        "druhy_ucastnik": u.get("Řidič – vozidlo B") or u.get("Účastník 2"),
        "vinik": u.get("Zavinění"),
        "zraneni": u.get("Zranění"),
        "poskozeni_vozidla": u.get("Poškození – vozidlo A"),
        "pojistitel_druheho": u.get("Pojistitel – vozidlo B"),
        "pojistna_smlouva_klienta": u.get("Pojistná smlouva – vozidlo A") or u.get("Pojistná smlouva"),
        "telefon_klienta": u.get("Telefon"),
        "vozidlo_klienta": u.get("Značka a model"),
    }
    return {k: v for k, v in f.items() if v}


def rule_naklad(t: str, u: dict) -> dict | None:
    if t == "faktura":
        return {"dodavatel": u.get("Dodavatel"), "cislo_dokladu": u.get("Číslo faktury"), "za_co": u.get("Za co"),
                "datum_vystaveni": u.get("Datum vystavení"), "splatnost": u.get("Splatnost"),
                "castka_celkem": u.get("_castka"), "mena": u.get("Měna", "CZK")}
    if t == "kalkulace":
        return {"dodavatel": u.get("Zpracoval"), "castka_celkem": u.get("_castka"), "mena": "CZK"}
    return None


# ----------------------------------------------------------------- zpracování souborů
def load_db() -> list[dict]:
    return json.loads(DB.read_text(encoding="utf-8")) if DB.exists() else []


def save_db(docs: list[dict]):
    DB.write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding="utf-8")


def free_name(folder: Path, name: str) -> Path:
    target, n = folder / name, 2
    while target.exists():
        target = folder / f"{Path(name).stem} ({n}){Path(name).suffix}"
        n += 1
    return target


def is_uninformative_filename(path: Path) -> bool:
    stem = path.stem.lower()
    if len(stem) < 7 or not stem.isalnum():
        return False
    letters = [char for char in stem if char.isalpha()]
    digits = [char for char in stem if char.isdigit()]
    no_vowels = not any(char in "aeiouy" for char in letters)
    high_entropy = len(set(stem)) / len(stem) >= 0.65
    return len(letters) >= 4 and len(digits) >= 2 and (no_vowels or high_entropy)


def descriptive_filename(doc: dict, original: Path) -> str:
    facts = doc.get("fakta") or {}
    expense = doc.get("naklad") or {}
    fields = doc.get("udaje") or {}
    doc_type = doc.get("typ")

    if doc_type == "faktura":
        identifier = expense.get("cislo_dokladu") or fields.get("Doklad") or fields.get("Číslo faktury")
        if identifier and not re.search(r"\d", str(identifier)):
            identifier = None
        subject = expense.get("za_co") or "doklad"
        parts = ["faktura", subject, identifier]
    elif doc_type == "protokol":
        parts = ["policejni_protokol", facts.get("datum_nehody"), facts.get("spz_klienta")]
    elif doc_type == "lekarska_zprava":
        parts = ["lekarska_zprava", fields.get("Datum vyšetření"), facts.get("spz_klienta")]
    elif doc_type == "doklad_vozidla":
        parts = ["doklad_vozidla", fields.get("Registrační značka") or facts.get("spz_klienta")]
    elif doc_type == "kalkulace":
        parts = ["kalkulace_opravy", fields.get("Zakázka") or facts.get("cislo_pojistne_udalosti")]
    elif doc_type == "hlaseni":
        parts = ["hlaseni_skody", facts.get("cislo_pojistne_udalosti") or facts.get("datum_nehody")]
    elif doc_type == "korespondence":
        parts = ["korespondence", fields.get("Číslo škody u druhé pojišťovny")]
    elif doc_type == "foto":
        parts = ["fotodokumentace", facts.get("spz_klienta") or "vozidla"]
    else:
        return original.name

    ascii_name = ascii_same_length("_".join(str(part) for part in parts if part))
    stem = re.sub(r"[^A-Za-z0-9]+", "_", ascii_name).strip("_").lower()[:90].rstrip("_")
    return f"{stem or 'dokument'}{original.suffix.lower()}"


def safe_read(path: Path) -> tuple[str, str]:
    try:
        return read_text(path)
    except Exception as e:  # poškozený soubor nesmí shodit celou aplikaci
        return "", f"chyba čtení: {e}"


class Zpracovani:
    def __init__(self):
        from ai_vytezeni import Analyzator, load_api_key
        key = load_api_key()
        self.ai = Analyzator(key) if key else None

    @property
    def rezim(self) -> str:
        from ai_vytezeni import MODEL
        if not self.ai:
            return "jednoduchá pravidla + OCR (bez AI)"
        if getattr(self.ai, "provider", None) == "local":
            return f"lokální model ({getattr(self.ai.local_client, 'model', 'llama3.1')})"
        return f"Claude ({MODEL})"

    def analyze_ai(self, path: Path, text: str, is_image: bool):
        try:
            return self.ai.analyze(path, text, is_image)
        except Exception as e:
            if getattr(self.ai, "provider", None) == "anthropic" and hasattr(self.ai, "anthropic"):
                auth = getattr(self.ai.anthropic, "AuthenticationError", None)
                conn = getattr(self.ai.anthropic, "APIConnectionError", None)
                if auth and isinstance(e, auth):
                    print("\n   ! Neplatný API klíč – dál pokračuji jen s pravidly. ", end="")
                    self.ai = None
                    return None
                if conn and isinstance(e, conn):
                    print("(bez připojení k internetu – použiji pravidla) ", end="")
                    return None
            print(f"(AI model selhal: {e} – použiji pravidla) ", end="")
        return None

    def process(self, path: Path) -> dict:
        is_image = path.suffix.lower() in IMAGE_EXT
        rec = {"soubor": path.name, "otisk": hashlib.sha1(path.read_bytes()).hexdigest(),
               "zpracovano": datetime.now().strftime("%d.%m.%Y %H:%M:%S"), "upozorneni": []}
        text, method = safe_read(path)
        doc_type, _ = classify(text, is_image)
        local_rules_sufficient = (getattr(self.ai, "provider", None) == "local"
                                  and doc_type not in {"foto", "neznamy"})
        a = self.analyze_ai(path, text, is_image) if self.ai and not local_rules_sufficient else None
        if a:
            zdroj = "lokální model" if getattr(self.ai, "provider", None) == "local" else "Claude"
            rec.update(
                typ="neznamy" if a.typ == "nesouvisejici" else a.typ, zdroj=zdroj,
                precteno="vyfocený/naskenovaný dokument" if a.vyfoceny_dokument else
                ("obrázek" if is_image else method or "PDF"),
                popis=a.popis, jazyk=a.jazyk, jistota=a.jistota, kvalita=a.kvalita, upozorneni=a.upozorneni,
                fakta={k: v for k, v in a.fakta.model_dump().items() if v},
                naklad=a.naklad.model_dump() if a.naklad else None,
                udaje={u.nazev: u.hodnota for u in a.udaje}, text=text.strip()[:4000])
        else:
            u = extract(doc_type, text, path.name) if doc_type != "neznamy" else {}
            rec.update(typ=doc_type, zdroj="pravidla", precteno=method, fakta=rule_fakta(doc_type, u),
                       naklad=rule_naklad(doc_type, u), udaje={k: v for k, v in u.items() if not k.startswith("_")},
                       text=text.strip()[:4000])
        folder = OUT / FOLDER[rec["typ"]]
        folder.mkdir(parents=True, exist_ok=True)
        output_name = descriptive_filename(rec, path) if is_uninformative_filename(path) else path.name
        target = free_name(folder, output_name)
        shutil.move(str(path), target)
        if target.name != path.name:
            rec["puvodni_soubor"] = path.name
        rec["cesta"] = str(target.relative_to(OUT))
        return rec


def ready_files() -> list[Path]:
    """Soubory, které se už dokopírovaly (velikost se 1 s nemění)."""
    files = [p for p in INBOX.iterdir() if p.is_file() and not p.name.startswith((".", "~$"))
             and p.suffix.lower() not in IGNORE_SUFFIX]
    sizes = {p: p.stat().st_size for p in files}
    time.sleep(1)
    return [p for p in files if p.exists() and p.stat().st_size == sizes[p]]


# ----------------------------------------------------------------- souhrn
FAKTA = [("cislo_pojistne_udalosti", "Číslo pojistné události"), ("klient", "Klient"),
         ("telefon_klienta", "Telefon klienta"), ("pojistna_smlouva_klienta", "Pojistná smlouva klienta"),
         ("vozidlo_klienta", "Vozidlo klienta"), ("spz_klienta", "SPZ vozidla klienta"), ("vin", "VIN"),
         ("datum_nehody", "Datum nehody"), ("cas_nehody", "Čas nehody"), ("misto_nehody", "Místo nehody"),
         ("vinik", "Viník"), ("druhy_ucastnik", "Druhý účastník"), ("spz_druheho_vozidla", "SPZ druhého vozidla"),
         ("pojistitel_druheho", "Pojistitel druhého účastníka"),
         ("cislo_skody_u_pojistitele_druheho", "Číslo škody u pojistitele druhého"),
         ("zraneni", "Zranění"), ("poskozeni_vozidla", "Poškození vozidla klienta")]
# U těchto údajů je rozdíl mezi dokumenty podezřelý (u jmen a míst jde často jen o jiný zápis).
MUSI_SEDET = ("spz_klienta", "vin", "datum_nehody", "spz_druheho_vozidla")
# Kterému dokumentu věřit víc, když údaj uvádí více dokumentů.
PRIORITA = ["protokol", "zaznam", "doklad_vozidla", "pojistna_smlouva", "ridicsky_prukaz", "hlaseni",
            "kalkulace", "faktura", "lekarska_zprava", "korespondence", "jiny_k_udalosti", "foto", "neznamy"]


def norm(v) -> str:
    return re.sub(r"[^a-z0-9]", "", ascii_same_length(str(v)).lower())


def money(v: float | None, mena: str | None = "CZK") -> str:
    if v is None:
        return "?"
    s = f"{v:,.2f}".replace(",", " ").replace(".", ",")
    return f"{s} Kč" if (mena or "CZK").upper() in ("CZK", "KČ", "KC") else f"{s} {mena}"


def case_overview(docs: list[dict]) -> tuple[list, list]:
    """Spojí údaje ze všech dokumentů do přehledu události + kontroly konzistence."""
    ordered = sorted(docs, key=lambda d: PRIORITA.index(d["typ"]) if d["typ"] in PRIORITA else 99)
    rows, checks = [], []
    for key, label in FAKTA:
        found = [(d["fakta"][key], d["soubor"]) for d in ordered if d.get("fakta", {}).get(key)]
        if not found:
            continue
        value, src = found[0]
        rows.append((label, value, src + (f" (+ {len(found) - 1} další)" if len(found) > 1 else "")))
        distinct = {}
        for v, s in found:
            distinct.setdefault(norm(v), (v, s))
        if key in MUSI_SEDET and len(distinct) > 1:
            listing = "; ".join(f"{v} ({s})" for v, s in distinct.values())
            checks.append(("warn", f"{label} se v dokumentech liší: {listing}"))
        elif key in MUSI_SEDET and len(found) > 1:
            checks.append(("ok", f"{label} je stejné ve {len(found)} dokumentech ({value})"))

    present = {d["typ"] for d in docs}
    for t in ("hlaseni", "doklad_vozidla", "foto", "kalkulace"):
        checks.append(("ok", f"{TYPE_LABEL[t]}: dodáno") if t in present else ("warn", f"{TYPE_LABEL[t]}: CHYBÍ"))
    checks.append(("ok", "Záznam o nehodě / policejní protokol: dodáno") if {"zaznam", "protokol"} & present
                  else ("warn", "Záznam o nehodě ani policejní protokol: CHYBÍ"))

    # Policie píše „bez zranění“, ale klient dodal lékařskou zprávu
    prot = next((d for d in docs if d["typ"] == "protokol" and d.get("fakta", {}).get("zraneni")), None)
    if prot and "lekarska_zprava" in present and re.search(r"\bbez\b|\bne\b|zadn", norm_spaces(prot["fakta"]["zraneni"])):
        checks.append(("warn", f"Policejní protokol uvádí „{prot['fakta']['zraneni']}“, ale je dodána lékařská "
                               "zpráva – ověřit nárok na zdravotní újmu"))

    # Kalkulace proti fakturám od stejného servisu
    for k in (d for d in docs if d["typ"] == "kalkulace" and (d.get("naklad") or {}).get("castka_celkem")):
        kn = k["naklad"]
        servis = norm(kn.get("dodavatel") or "")[:8]
        fakt = [d for d in docs if d["typ"] == "faktura" and (d.get("naklad") or {}).get("castka_celkem")
                and servis and norm(d["naklad"].get("dodavatel") or "").startswith(servis)]
        for d in fakt:
            diff = d["naklad"]["castka_celkem"] - kn["castka_celkem"]
            checks.append(("ok", f"Faktura {d['soubor']} ({money(d['naklad']['castka_celkem'])}) odpovídá kalkulaci")
                          if abs(diff) < 1 else
                          ("warn", f"Faktura {d['soubor']} se liší od kalkulace o {money(diff)}"))

    seen = {}
    for d in docs:
        if d.get("otisk") in seen:
            checks.append(("warn", f"{d['soubor']} je stejný soubor jako {seen[d['otisk']]} (nahráno dvakrát)"))
        seen.setdefault(d.get("otisk"), d["soubor"])
        if d.get("jistota") == "nizka":
            why = (d.get("kvalita") or "údaje jsou nejisté").rstrip(". ")
            checks.append(("warn", f"{d['soubor']}: špatně čitelné ({why}) – vyžádat lepší kopii"))
        for u in d.get("upozorneni") or []:
            checks.append(("warn", f"{d['soubor']}: {u}"))
        if d["typ"] == "neznamy":
            checks.append(("warn", f"{d['soubor']}: nerozpoznáno / nesouvisí s událostí – zkontrolujte ručně"))
    return rows, checks


def norm_spaces(v: str) -> str:
    return ascii_same_length(v).lower()


def write_summary(docs: list[dict], rezim: str):
    e = html.escape
    rows, checks = case_overview(docs)
    out = [f"<!doctype html><html lang='cs'><head><meta charset='utf-8'>"
           f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
           f"<title>Souhrn pojistné události</title><style>{CSS}</style></head><body><main>",
           "<h1>Souhrn pojistné události</h1>",
           f"<p class='muted'>Vygenerováno automaticky {datetime.now():%d.%m.%Y %H:%M} z {len(docs)} dokumentů · "
           f"rozpoznávání: {e(rezim)}. Údaje jsou vytěžené strojově – před výplatou je ověřte.</p>"]

    out.append("<h2>Přehled události</h2><div class='card'><table>")
    out += [f"<tr><th>{e(k)}</th><td>{e(str(v))}<div class='src'>zdroj: {e(s)}</div></td></tr>" for k, v, s in rows]
    out.append("</table></div>")

    out.append("<h2>Kontroly a chybějící podklady</h2><div class='card'>")
    out += [f"<div class='check {c}'>{'✔' if c == 'ok' else '⚠'} {e(t)}</div>"
            for c, t in sorted(checks, key=lambda c: c[0] == "ok")]
    out.append("</div>")

    costs = [d for d in docs if d["typ"] == "faktura" and d.get("naklad")]
    if costs:
        out.append("<h2>Náklady (faktury a účtenky)</h2><div class='card'><table><tr><th>Za co</th>"
                   "<th>Dodavatel</th><th>Doklad / splatnost</th><th class='num'>Částka</th></tr>")
        totals: dict[str, float] = {}
        for d in costs:
            n = d["naklad"]
            mena = (n.get("mena") or "CZK").upper()
            if n.get("castka_celkem") is not None:
                totals[mena] = totals.get(mena, 0) + n["castka_celkem"]
            due = "uhrazeno" if n.get("uhrazeno") else f"splatnost {n.get('splatnost') or '–'}"
            out.append(f"<tr><td>{e(n.get('za_co') or '')}</td><td>{e(n.get('dodavatel') or '')}</td>"
                       f"<td>{e(n.get('cislo_dokladu') or '?')}<br><span class='src'>{e(due)}</span></td>"
                       f"<td class='num'>{e(money(n.get('castka_celkem'), mena))}</td></tr>")
        for mena, total in totals.items():
            out.append(f"<tr><td colspan='3'><b>Celkem {e(mena)}</b></td>"
                       f"<td class='num'><b>{e(money(total, mena))}</b></td></tr>")
        out.append("</table></div>")

    photos = [d for d in docs if d["typ"] == "foto"]
    if photos:
        out.append("<h2>Fotodokumentace</h2><div class='card photos'>")
        out += [f"<figure><a href='{e(d['cesta'])}'><img src='{e(d['cesta'])}' alt='{e(d['soubor'])}'></a>"
                f"<figcaption>{e(d.get('popis') or d['soubor'])}</figcaption></figure>" for d in photos]
        out.append("</div>")

    out.append("<h2>Jednotlivé dokumenty</h2>")
    for d in sorted(docs, key=lambda d: PRIORITA.index(d["typ"]) if d["typ"] in PRIORITA else 99):
        tags = [d.get("zdroj", ""), d.get("precteno", "")]
        if d.get("jistota"):
            tags.append(f"jistota: {d['jistota']}")
        if d.get("jazyk") and d["jazyk"].lower() not in ("čeština", "cestina"):
            tags.append(d["jazyk"])
        out.append(f"<div class='card'><b>{e(TYPE_LABEL[d['typ']])}</b>"
                   + "".join(f"<span class='tag'>{e(t)}</span>" for t in tags if t)
                   + f"<div class='src'><a href='{e(d['cesta'])}'>{e(d['cesta'])}</a> · zpracováno "
                     f"{e(d['zpracovano'])}</div>")
        if d.get("puvodni_soubor"):
            out.append(f"<div class='src'>Původní název uploadu: {e(d['puvodni_soubor'])}</div>")
        if d.get("popis"):
            out.append(f"<p>{e(d['popis'])}</p>")
        if d.get("kvalita"):
            out.append(f"<div class='check warn'>⚠ Čitelnost: {e(d['kvalita'])}</div>")
        if d.get("udaje"):
            out.append("<table>" + "".join(f"<tr><th>{e(k)}</th><td>{e(str(v))}</td></tr>"
                                           for k, v in d["udaje"].items()) + "</table>")
        elif d["typ"] == "neznamy":
            out.append("<p class='muted'>Typ dokumentu nebyl rozpoznán.</p>")
        if d.get("fakta"):
            labels = dict(FAKTA)
            out.append("<details><summary>Údaje o nehodě z tohoto dokumentu</summary><table>"
                       + "".join(f"<tr><th>{e(labels.get(k, k))}</th><td>{e(str(v))}</td></tr>"
                                 for k, v in d["fakta"].items()) + "</table></details>")
        if d.get("text"):
            out.append(f"<details><summary>Přečtený text</summary><pre>{e(d['text'])}</pre></details>")
        out.append("</div>")
    out.append("</main></body></html>")
    SUMMARY.write_text("\n".join(out), encoding="utf-8")


# ----------------------------------------------------------------- hlavní smyčka
def run_once(z: Zpracovani, docs: list[dict]) -> bool:
    files = ready_files()
    by_hash = {d["otisk"]: d for d in docs if d.get("otisk")}
    for p in files:
        print(f"→ {p.name} ... ", end="", flush=True)
        try:
            digest = hashlib.sha1(p.read_bytes()).hexdigest()
        except FileNotFoundError:
            print("soubor už mezitím zmizel, přeskakuji.")
            continue
        previous = by_hash.get(digest)
        if previous:
            archived = OUT / (previous.get("cesta") or "")
            if archived.is_file():
                warning = f"Duplicitní nahrání přeskočeno; stejný soubor už je evidován jako {previous['soubor']}."
                warnings = previous.setdefault("upozorneni", [])
                if warnings is None:
                    warnings = previous["upozorneni"] = []
                if warning not in warnings:
                    warnings.append(warning)
                p.unlink(missing_ok=True)
                save_db(docs)
                print(f"duplicitní obsah už je zařazen jako {previous['soubor']}; kopie přeskočena.")
                continue
            print(f"chybí archivní soubor {previous.get('cesta')}; obnovuji znovu nahranou kopii. ", end="")
        try:
            doc = z.process(p)
        except FileNotFoundError:
            if not p.exists():
                print("soubor už mezitím přesunul jiný běh aplikace, přeskakuji.")
                continue
            raise
        if previous:
            doc["upozorneni"] = list(previous.get("upozorneni") or [])
            docs[next(i for i, item in enumerate(docs) if item is previous)] = doc
        else:
            docs.append(doc)
        if doc.get("otisk"):
            by_hash[doc["otisk"]] = doc
        save_db(docs)
        extra = f", jistota {doc['jistota']}" if doc.get("jistota") else ""
        print(f"{TYPE_LABEL[doc['typ']]}  ({doc['zdroj']}{extra})  →  roztridene/{doc['cesta']}")
    if files:
        write_summary(docs, z.rezim)
        print(f"   Souhrn aktualizován: {SUMMARY}")
    return bool(files)


def main():
    sys.stdout.reconfigure(line_buffering=True)
    INBOX.mkdir(exist_ok=True)
    OUT.mkdir(exist_ok=True)
    docs = load_db()
    z = Zpracovani()
    print(f"Rozpoznávání: {z.rezim}")
    if "--jednou" in sys.argv:
        if not run_once(z, docs):
            print(f"Ve složce {INBOX} nejsou žádné soubory.")
        return
    print(f"Hlídám složku: {INBOX}\nVložte do ní dokumenty. Ukončení: Ctrl+C\n")
    try:
        while True:
            run_once(z, docs)
            time.sleep(1)
    except KeyboardInterrupt:
        print("\nKonec.")


if __name__ == "__main__":
    main()
